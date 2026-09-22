import json,csv,collections,copy,openpyxl
from openpyxl.styles import Font
exec(open('build1.py').read().split('# geonames cities')[0])
U='inputs/'
O='../dataset/'
st=json.load(open('states_m.json')); cities=json.load(open('cities_m.json'))
A=json.load(open('areas.json'))['admin1']; use=collections.Counter(f for s in st for f in s['_f'])
wds=json.load(open('wd_state_area.json')); ca=json.load(open('wd_a.json'))['city_area']
# ---- state areas
sarea={}
for s in st:
    k=(s['country'],s['state_province'])
    if s['area_km2']: sarea[k]=(float(s['area_km2']),'Natural Earth (calculated)')
    elif s['_f'] and all(use[f]==1 for f in s['_f']) and sum(A[f] for f in s['_f'])>0.5: sarea[k]=(round(sum(A[f] for f in s['_f']),1),'Natural Earth (calculated)')
    elif '|'.join(k) in wds: sarea[k]=(wds['|'.join(k)],'Wikidata')
# ---- city areas (Wikidata, sanity filtered)
carea={}; dropped=0
for i,c in enumerate(cities):
    a=ca.get(str(c['_g']['id']))
    if a is None: continue
    d=int(c['total_population'])/a if a>0 else 0
    if a>=0.5 and 1<=d<=60000: carea[i]=a
    else: dropped+=1
print('state areas',len(sarea),'/',len(st),' city areas',len(carea),'dropped',dropped)
# ---- workbook
wb=openpyxl.load_workbook(U+'e9c447d7-World_Population_by_Continent_Country_State_City_1.xlsx')
# countries area from workbook
cw={r[1]:(r[17],r[18]) for r in wb['Countries'].iter_rows(min_row=2,values_only=True) if r[1]}
ws=wb['States_Provinces']; blue=copy.copy(ws['L2'].font); hdr=ws['L1']
ws['N1']='Area Source'; ws['N1'].font=copy.copy(hdr.font); ws['N1'].fill=copy.copy(hdr.fill); ws['N1'].alignment=copy.copy(hdr.alignment); ws['N1'].border=copy.copy(hdr.border)
ws.column_dimensions['N'].width=26; filled=0
for r in range(2,ws.max_row+1):
    k=(ws.cell(r,2).value,ws.cell(r,3).value)
    if k in sarea:
        if ws.cell(r,12).value in (None,''):
            ws.cell(r,12).value=sarea[k][0]; ws.cell(r,12).font=copy.copy(blue); ws.cell(r,12).number_format=ws['L2'].number_format; filled+=1
        ws.cell(r,14).value=sarea[k][1]; ws.cell(r,14).font=copy.copy(blue)
    if ws.cell(r,13).value in (None,''): ws.cell(r,13).value=f'=IF(L{r}="","",F{r}/L{r})'; ws.cell(r,13).number_format=ws['M2'].number_format
if ws.auto_filter.ref: ws.auto_filter.ref=f'A1:N{ws.max_row}'
print('workbook states filled',filled)
ws=wb['Cities']; hdr=ws['O1']; cfont=copy.copy(ws['F2'].font)
for col,name,w in [('P','Area (km²)',12),('Q','Population Density (per km²)',16),('R','Area Source',12)]:
    c=ws[col+'1']; c.value=name; c.font=copy.copy(hdr.font); c.fill=copy.copy(hdr.fill); c.alignment=copy.copy(hdr.alignment); c.border=copy.copy(hdr.border); ws.column_dimensions[col].width=w
mism=0
for r in range(2,ws.max_row+1):
    i=r-2; c=cities[i]
    if (ws.cell(r,2).value,ws.cell(r,5).value)!=(c['country'],c['city']): mism+=1; continue
    if i in carea:
        ws.cell(r,16).value=round(carea[i],2); ws.cell(r,16).font=copy.copy(cfont); ws.cell(r,16).number_format='#,##0.0'
        ws.cell(r,18).value='Wikidata'; ws.cell(r,18).font=copy.copy(cfont)
    ws.cell(r,17).value=f'=IF(P{r}="","",F{r}/P{r})'; ws.cell(r,17).number_format='#,##0.0'
ws.auto_filter.ref=f'A1:R{ws.max_row}'; print('city row mismatches',mism)
rm=wb['Read Me']
for row in rm.iter_rows():
    for c in row:
        v=c.value or ''
        if isinstance(v,str) and v.startswith('City area: not included'):
            c.value=('City area: Wikidata "area" (P2046), linked by GeoNames ID, for about '+f'{len(carea):,}'+' of 34,143 cities. '
                     'Areas are the administrative city limits Wikidata records; values implying fewer than 1 or more than 60,000 people per km² were dropped as likely errors. Blank = not recorded.')
        if isinstance(v,str) and v.startswith('State/province area: calculated'):
            c.value=(c.value+' Gaps are now filled where possible: more Natural Earth shapes matched by GeoNames code, then Wikidata areas. The Area Source column says which.')
wb.save(O+'World_Population_by_Continent_Country_State_City.xlsx')
# ---- CSVs
def wcsv(name,header,rows):
    with open(O+name,'w',newline='',encoding='utf8') as f: w=csv.writer(f); w.writerow(header); w.writerows(rows)
def rd(p): return list(csv.DictReader(open(U+p,encoding='utf8')))
C=rd('b7c56634-countries.csv'); rows=[]
for c in C:
    s,l=cw.get(c['country'],(None,None)); p=int(c['total_population'])
    rows.append(list(c.values())+[s if s is not None else '', l if l is not None else '', round(p/l,2) if l else ''])
wcsv('countries.csv',list(C[0].keys())+['surface_area_km2','land_area_km2','pop_density_per_km2'],rows)
K=rd('038bfaee-continents.csv'); rows=[]
for k in K:
    cs=[c for c in C if c['continent']==k['continent']]
    s=sum(float(cw[c['country']][0] or 0) for c in cs); l=sum(float(cw[c['country']][1] or 0) for c in cs)
    rows.append(list(k.values())+[round(s,1),round(l,1),round(int(k['total_population'])/l,2) if l else ''])
wcsv('continents.csv',list(K[0].keys())+['surface_area_km2','land_area_km2','pop_density_per_km2'],rows)
S=rd('07e48944-states_provinces.csv'); rows=[]
for s in S:
    k=(s['country'],s['state_province']); a,src=sarea.get(k,('',''))
    s['area_km2']=a; s['city_pop_density_per_km2']=round(int(s['population_in_listed_cities'])/a,1) if a else ''
    rows.append(list(s.values())+[src])
wcsv('states_provinces.csv',list(S[0].keys())+['area_source'],rows)
T=rd('d1862297-cities.csv'); rows=[]
for i,c in enumerate(T):
    a=carea.get(i); rows.append(list(c.values())+[round(a,2) if a else '', round(int(c['total_population'])/a,1) if a else '', 'Wikidata' if a else ''])
wcsv('cities.csv',list(T[0].keys())+['area_km2','pop_density_per_km2','area_source'],rows)
# schemas
def sch(src,extra):
    s=json.load(open(U+src)); names={x['name'] for x in s}
    s+= [{'name':n,'type':t,'mode':'NULLABLE'} for n,t in extra if n not in names]; return s
json.dump(sch('0b1196ce-countries_schema.json',[('surface_area_km2','FLOAT'),('land_area_km2','FLOAT'),('pop_density_per_km2','FLOAT')]),open(O+'countries_schema.json','w'),indent=2)
json.dump(sch('213f3126-continents_schema.json',[('surface_area_km2','FLOAT'),('land_area_km2','FLOAT'),('pop_density_per_km2','FLOAT')]),open(O+'continents_schema.json','w'),indent=2)
json.dump(sch('faaf4a1e-states_provinces_schema.json',[('area_source','STRING')]),open(O+'states_provinces_schema.json','w'),indent=2)
json.dump(sch('f5475a29-cities_schema.json',[('area_km2','FLOAT'),('pop_density_per_km2','FLOAT'),('area_source','STRING')]),open(O+'cities_schema.json','w'),indent=2)
import shutil
shutil.copy(U+'30cdbdec-geo_features.csv',O+'geo_features.csv'); shutil.copy(U+'ac5ea00f-geo_features_schema.json',O+'geo_features_schema.json')
json.dump({'sarea':{'|'.join(k):v for k,v in sarea.items()},'carea':{str(i):a for i,a in carea.items()},'cw':cw},open('area_final.json','w'))
